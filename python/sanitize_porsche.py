#!/usr/bin/env python3
"""Agente de higienização de vendas Porsche.

Lê a planilha bruta, aplica as regras do Schema.md e grava uma nova planilha
com as colunas higienizadas logo após cada coluna de origem.

Uso:
    python3 sanitize_porsche.py ENTRADA.xlsx SAIDA.xlsx [--sheet NOME]
"""
import argparse
import datetime as dt
import re
import sys
import unicodedata
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

INV = "INVALID"
INV_PT = "INVÁLIDO"


# ----------------------------------------------------------------- utilidades
def strip_accents(s):
    return "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))


def norm(s):
    """minúsculas, sem acento, espaços colapsados."""
    return re.sub(r"\s+", " ", strip_accents(str(s)).lower()).strip()


def blank(v):
    return v is None or (isinstance(v, str) and not v.strip())


def title_words(s):
    out = []
    for w in re.sub(r"\s+", " ", str(s)).strip().split(" "):
        if any(ch.isdigit() for ch in w):
            out.append(w.upper())
        else:
            out.append("-".join(p[:1].upper() + p[1:].lower() for p in w.split("-")))
    return " ".join(out)


# ------------------------------------------------------- números por extenso
WORDS = {
    # inglês
    "zero": 0, "one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6, "seven": 7,
    "eight": 8, "nine": 9, "ten": 10, "eleven": 11, "twelve": 12, "thirteen": 13,
    "fourteen": 14, "fifteen": 15, "sixteen": 16, "seventeen": 17, "eighteen": 18,
    "nineteen": 19, "twenty": 20, "thirty": 30, "forty": 40, "fifty": 50, "sixty": 60,
    "seventy": 70, "eighty": 80, "ninety": 90,
    # português
    "um": 1, "uma": 1, "dois": 2, "duas": 2, "tres": 3, "quatro": 4, "cinco": 5, "seis": 6,
    "sete": 7, "oito": 8, "nove": 9, "dez": 10, "onze": 11, "doze": 12, "treze": 13,
    "quatorze": 14, "catorze": 14, "quinze": 15, "dezesseis": 16, "dezessete": 17,
    "dezoito": 18, "dezenove": 19, "vinte": 20, "trinta": 30, "quarenta": 40,
    "cinquenta": 50, "sessenta": 60, "setenta": 70, "oitenta": 80, "noventa": 90,
    "cem": 100, "cento": 100, "duzentos": 200, "trezentos": 300, "quatrocentos": 400,
    "quinhentos": 500, "seiscentos": 600, "setecentos": 700, "oitocentos": 800,
    "novecentos": 900,
}
FILLER = {"and", "e", "usd", "dollars", "dollar", "dolares", "dolar", "americanos",
          "americano", "us", "miles", "mile", "milhas", "milha", "mi", "de", "reais"}


def words_to_number(text):
    """Converte número por extenso (EN/PT) em int; None se não for possível."""
    toks = [t for t in re.split(r"[\s\-]+", norm(text)) if t and t not in FILLER]
    if not toks:
        return None
    total = cur = 0
    for t in toks:
        if t in WORDS:
            cur += WORDS[t]
        elif t == "hundred":
            cur = max(cur, 1) * 100
        elif t in ("thousand", "mil"):
            total += max(cur, 1) * 1000
            cur = 0
        elif t in ("million", "milhao", "milhoes"):
            total += max(cur, 1) * 1_000_000
            cur = 0
        else:
            return None
    return total + cur


def words_to_year(text):
    """'twenty twenty four' -> 2024; 'two thousand twenty one' -> 2021."""
    toks = [t for t in re.split(r"[\s\-]+", norm(text)) if t and t not in ("and", "e")]
    if not toks:
        return None
    n = words_to_number(" ".join(toks))
    if n is not None and n >= 1000:
        return n
    for i in range(1, len(toks)):  # forma "vinte vinte e quatro"
        left, right = words_to_number(" ".join(toks[:i])), words_to_number(" ".join(toks[i:]))
        if left in (19, 20) and right is not None and 0 <= right < 100:
            return left * 100 + right
    return None


# ------------------------------------------------------------ números em dígitos
def parse_amount(numstr):
    """Interpreta '1.200', '19,250', '89.750,00', '9.5' -> Decimal."""
    s = numstr.strip().strip(".,")
    has_dot, has_com = "." in s, "," in s
    if has_dot and has_com:
        dec = "." if s.rfind(".") > s.rfind(",") else ","
        s = s.replace("," if dec == "." else ".", "").replace(dec, ".")
    elif has_dot or has_com:
        sep = "." if has_dot else ","
        parts = s.split(sep)
        if len(parts) > 2 or (len(parts) == 2 and len(parts[1]) == 3):
            s = "".join(parts)  # separador de milhar
        else:
            s = s.replace(sep, ".")  # separador decimal
    return Decimal(s)


NUM_RE = re.compile(r"\d[\d.,]*")


# --------------------------------------------------------------------- datas
MONTHS = ["january", "february", "march", "april", "may", "june", "july", "august",
          "september", "october", "november", "december"]
MONTHS_PT = ["janeiro", "fevereiro", "marco", "abril", "maio", "junho", "julho", "agosto",
             "setembro", "outubro", "novembro", "dezembro"]


def month_from_name(word):
    w = norm(word).rstrip(".")
    if len(w) < 3:
        return None
    for names in (MONTHS, MONTHS_PT):
        for i, m in enumerate(names, 1):
            if m.startswith(w) and w[:3] == m[:3]:
                return i
    return None


def sanitize_date(v):
    if blank(v):
        return INV_PT
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = str(v).strip()
    y = m = d = None
    if (g := re.fullmatch(r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", s)):
        y, m, d = map(int, g.groups())
    elif (g := re.fullmatch(r"(\d{1,2})[-/](\d{1,2})[-/](\d{4}|\d{2})", s)):
        m, d, y = int(g[1]), int(g[2]), int(g[3])
        if y < 100:
            y += 2000
    elif (g := re.fullmatch(r"(?:[A-Za-z]{3,9}\.?,?\s+)?([A-Za-z]+)\.?\s+(\d{1,2})(?:st|nd|rd|th)?,?\s+(\d{4})", s)):
        m, d, y = month_from_name(g[1]), int(g[2]), int(g[3])
    elif (g := re.fullmatch(r"(\d{1,2})º?\s+de\s+([A-Za-zçÇ]+)\s+de\s+(\d{4})", s, re.I)):
        d, m, y = int(g[1]), month_from_name(g[2]), int(g[3])
    if not (y and m and d):
        return INV_PT
    try:
        return dt.date(y, m, d).isoformat()
    except ValueError:
        return INV_PT


# -------------------------------------------------------------------- modelos
CANON_MODELS = [
    "911 Carrera", "911 Carrera S", "911 Carrera GTS", "911 Turbo", "911 Turbo S", "911 GT3",
    "911 GT3 RS", "911 Dakar", "911 Targa 4", "911 Targa 4S", "718 Cayman", "718 Cayman S",
    "718 Cayman GT4 RS", "718 Boxster", "718 Boxster GTS", "718 Spyder RS",
    "Cayenne",  # o Schema traz "Pimenta" (tradução automática de "Cayenne")
    "Cayenne S", "Cayenne Coupé", "Cayenne E-Hybrid", "Cayenne Turbo", "Cayenne Turbo GT",
    "Macan", "Macan S", "Macan T", "Macan GTS", "Macan Electric", "Panamera", "Panamera 4",
    "Panamera 4S", "Panamera Turbo", "Panamera Turbo S", "Panamera 4 E-Hybrid", "Taycan",
    "Taycan 4S", "Taycan GTS", "Taycan Turbo", "Taycan Turbo S", "Taycan Cross Turismo",
]
MODEL_MAP = {norm(m): m for m in CANON_MODELS}
MODEL_MAP.update({"pimenta": "Cayenne", "macan eletrico": "Macan Electric",
                  "cayenne coupe": "Cayenne Coupé"})


def sanitize_model(v):
    if blank(v):
        return INV
    key = norm(v)
    return MODEL_MAP.get(key) or title_words(v)


# ------------------------------------------------------------------ ano modelo
def sanitize_year(v):
    if blank(v):
        return INV
    y = None
    if isinstance(v, (int, float)) and float(v).is_integer():
        y = int(v)
    else:
        s = str(v).strip()
        if re.fullmatch(r"\d{4}", s):
            y = int(s)
        elif (g := re.fullmatch(r"(\d{2})[-\s](\d{2})", s)):
            y = int(g[1] + g[2])
        else:
            y = words_to_year(s)
    return str(y) if y is not None and 1990 <= y <= 2035 else INV


# ----------------------------------------------------------------------- preço
def sanitize_price(v):
    if blank(v):
        return INV
    try:
        if isinstance(v, (int, float)):
            amt = Decimal(str(v))
        else:
            s = str(v).strip()
            m = NUM_RE.search(s)
            if m:
                amt = parse_amount(m.group())
                if re.search(r"\d\s*k\b", s, re.I):
                    amt *= 1000
                elif re.search(r"\bmil\b", s, re.I):
                    amt *= 1000
                elif re.search(r"\bmilh(ao|oes|ão|ões)\b|\bmillion\b", s, re.I):
                    amt *= 1_000_000
            else:
                n = words_to_number(s)
                if n is None:
                    return INV
                amt = Decimal(n)
        if amt <= 0:
            return INV
        return f"{amt.quantize(Decimal('0.01'), ROUND_HALF_UP):.2f}"
    except (InvalidOperation, ValueError):
        return INV


# ---------------------------------------------------------------- quilometragem
NEW_WORDS = {"new", "novo", "nova", "zero"}
KM_RE = re.compile(r"\bkm\b|\bkms\b|quilometro|kilometer", re.I)


def sanitize_mileage(v):
    if blank(v):
        return INV
    try:
        if isinstance(v, (int, float)):
            amt = Decimal(str(v))
        else:
            s = str(v).strip()
            m = NUM_RE.search(s)
            if m:
                amt = parse_amount(m.group())
                if KM_RE.search(strip_accents(s)):
                    amt *= Decimal("0.621371")
            else:
                toks = set(norm(s).replace("-", " ").split())
                if toks & NEW_WORDS and toks <= NEW_WORDS | {"car", "carro", "miles", "milhas", "mi"}:
                    return 0
                n = words_to_number(s)
                if n is None:
                    return INV
                amt = Decimal(n)
        if amt < 0:
            return INV
        return int(amt.quantize(Decimal("1"), ROUND_HALF_UP))
    except (InvalidOperation, ValueError):
        return INV


# ------------------------------------------------------------------- pagamento
PAYMENT_MAP = {
    "credit card": "Cartão de Crédito", "creditcard": "Cartão de Crédito",
    "credit": "Cartão de Crédito", "credit card payment": "Cartão de Crédito",
    "cartao de credito": "Cartão de Crédito",
    "debit card": "Cartão de Débito", "debit": "Cartão de Débito",
    "cartao de debito": "Cartão de Débito",
    "bank transfer": "Transferência Bancária", "banktransfer": "Transferência Bancária",
    "bank wire": "Transferência Bancária", "wire transfer": "Transferência Bancária",
    "wiretransfer": "Transferência Bancária", "wire": "Transferência Bancária",
    "transferencia bancaria": "Transferência Bancária", "transferencia": "Transferência Bancária",
    "financing": "Financiamento", "financing plan": "Financiamento", "finance": "Financiamento",
    "financiamento": "Financiamento",
    "lease": "Arrendamento", "leasing": "Arrendamento", "lease plan": "Arrendamento",
    "arrendamento": "Arrendamento",
    "cash": "Dinheiro", "cash payment": "Dinheiro", "dinheiro": "Dinheiro",
    "ach": "Pagamento ACH", "ach payment": "Pagamento ACH", "pagamento ach": "Pagamento ACH",
    "crypto": "Pagamento em Criptomoedas", "crypto payment": "Pagamento em Criptomoedas",
    "cryptocurrency": "Pagamento em Criptomoedas", "criptomoeda": "Pagamento em Criptomoedas",
    "criptomoedas": "Pagamento em Criptomoedas",
    "pagamento em criptomoedas": "Pagamento em Criptomoedas",
}


def sanitize_payment(v):
    if blank(v):
        return INV
    key = norm(re.sub(r"[-_]+", " ", str(v)))
    return PAYMENT_MAP.get(key) or PAYMENT_MAP.get(key.replace(" ", "")) or title_words(v)


# ---------------------------------------------------------------------- cidade
def sanitize_city(v):
    if blank(v):
        return INV
    out = []
    for w in re.sub(r"\s+", " ", str(v)).strip().split(" "):
        out.append("-".join(p[:1].upper() + p[1:].lower() for p in w.split("-")))
    return " ".join(out)


# ---------------------------------------------------------------------- estado
STATES = {
    "AL": "Alabama", "AK": "Alaska", "AZ": "Arizona", "AR": "Arkansas", "CA": "California",
    "CO": "Colorado", "CT": "Connecticut", "DE": "Delaware", "DC": "District of Columbia",
    "FL": "Florida", "GA": "Georgia", "HI": "Hawaii", "ID": "Idaho", "IL": "Illinois",
    "IN": "Indiana", "IA": "Iowa", "KS": "Kansas", "KY": "Kentucky", "LA": "Louisiana",
    "ME": "Maine", "MD": "Maryland", "MA": "Massachusetts", "MI": "Michigan",
    "MN": "Minnesota", "MS": "Mississippi", "MO": "Missouri", "MT": "Montana",
    "NE": "Nebraska", "NV": "Nevada", "NH": "New Hampshire", "NJ": "New Jersey",
    "NM": "New Mexico", "NY": "New York", "NC": "North Carolina", "ND": "North Dakota",
    "OH": "Ohio", "OK": "Oklahoma", "OR": "Oregon", "PA": "Pennsylvania",
    "RI": "Rhode Island", "SC": "South Carolina", "SD": "South Dakota", "TN": "Tennessee",
    "TX": "Texas", "UT": "Utah", "VT": "Vermont", "VA": "Virginia", "WA": "Washington",
    "WV": "West Virginia", "WI": "Wisconsin", "WY": "Wyoming",
}
STATE_MAP = {norm(n): c for c, n in STATES.items()}
STATE_MAP.update({  # nomes em português
    "nova iorque": "NY", "nova york": "NY", "florida": "FL", "georgia": "GA", "havai": "HI",
    "luisiana": "LA", "nova jersey": "NJ", "nova hampshire": "NH", "novo mexico": "NM",
    "carolina do norte": "NC", "carolina do sul": "SC", "dakota do norte": "ND",
    "dakota do sul": "SD", "virginia ocidental": "WV", "pensilvania": "PA",
    "misuri": "MO", "mississipi": "MS", "washington dc": "DC", "distrito de columbia": "DC",
})


def sanitize_state(v):
    if blank(v):
        return INV_PT
    s = re.sub(r"\.", "", str(v)).strip()
    if s.upper() in STATES:
        return s.upper()
    return STATE_MAP.get(norm(s), INV_PT)


# ---------------------------------------------------------------------- status
STATUS_MAP = {
    "delivered": "Entregue", "deliverd": "Entregue", "entregue": "Entregue",
    "entregado": "Entregue",
    "pending": "Pendente", "pendente": "Pendente",
    "in transit": "Em trânsito", "intransit": "Em trânsito", "em transito": "Em trânsito",
    "cancelled": "Cancelado", "canceled": "Cancelado", "cancelado": "Cancelado",
    "awaiting delivery": "Aguardando Entrega", "aguardando entrega": "Aguardando Entrega",
    "awaiting pickup": "Aguardando Recolhimento",
    "aguardando recolhimento": "Aguardando Recolhimento",
    "pending approval": "Aguardando aprovação", "aguardando aprovacao": "Aguardando aprovação",
    "pending review": "Em análise", "under review": "Em análise", "em analise": "Em análise",
    "shipped": "Enviado", "enviado": "Enviado",
    "awaiting review": "Aguardando revisão", "aguardando revisao": "Aguardando revisão",
}


def sanitize_status(v):
    if blank(v):
        return INV
    key = norm(re.sub(r"[-_]+", " ", re.sub(r"[^\w\s\-]", "", str(v))))
    return STATUS_MAP.get(key, INV)


# ------------------------------------------------------------- mapa de colunas
# (coluna de origem na planilha, coluna de origem no Schema, saída, função)
RULES = [
    ("sale_date", "data_da_venda", "DataDaVendaSanitizada", sanitize_date),
    ("porsche_model", "porsche_model", "PorscheModelSanitized", sanitize_model),
    ("model_year", "model_year", "ModelYearSanitized", sanitize_year),
    ("sale_price", "preço_de_venda", "Preço_de_vendaSanitizado", sanitize_price),
    ("vehicle_mileage", "vehicle_mileage", "VehicleMileageSanitized", sanitize_mileage),
    ("payment_method", "método_de_pagamento", "Método_de_pagamentoSanitizado", sanitize_payment),
    ("city", "cidade", "CidadeSanitizada", sanitize_city),
    ("state", "estado", "EstadoSanitizado", sanitize_state),
    ("delivery_status", "status_de_entrega", "Status_de_entrega_higienizado", sanitize_status),
]
# aliases para cabeçalhos em português do Schema
ALIASES = {"data_da_venda": "sale_date", "preço_de_venda": "sale_price",
           "método_de_pagamento": "payment_method", "cidade": "city", "estado": "state",
           "status_da_entrega": "delivery_status", "status_de_entrega": "delivery_status"}


def is_invalid(x):
    return x in (INV, INV_PT)


def run(src, dst, sheet=None):
    wb = openpyxl.load_workbook(src)
    ws = wb[sheet] if sheet else wb.active
    header = [c.value for c in ws[1]]
    canon = [ALIASES.get(h, h) for h in header]
    rules = {r[0]: r for r in RULES}

    out_wb = openpyxl.Workbook()
    out = out_wb.active
    out.title = ws.title[:31]
    head_fill = PatternFill("solid", fgColor="1F2937")
    new_fill = PatternFill("solid", fgColor="0F766E")
    bad_fill = PatternFill("solid", fgColor="FECACA")

    # colunas de saída: cada coluna sanitizada imediatamente após a origem
    layout = []  # (índice_origem, tipo, regra)
    for i, c in enumerate(canon):
        layout.append((i, "orig", None))
        if c in rules:
            layout.append((i, "san", rules[c]))
    for j, (i, kind, rule) in enumerate(layout, 1):
        cell = out.cell(1, j, header[i] if kind == "orig" else rule[2])
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill if kind == "orig" else new_fill

    stats = {r[2]: 0 for r in RULES}
    n = 0
    for r_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if all(blank(v) for v in row):
            continue
        n += 1
        for j, (i, kind, rule) in enumerate(layout, 1):
            raw = row[i]
            if kind == "orig":
                cell = out.cell(r_idx, j, raw)
                if isinstance(raw, (dt.datetime, dt.date)):
                    cell.number_format = "yyyy-mm-dd"
            else:
                val = rule[3](raw)
                cell = out.cell(r_idx, j, val)
                if is_invalid(val):
                    cell.fill = bad_fill
                    stats[rule[2]] += 1
                if rule[2] == "Preço_de_vendaSanitizado":
                    cell.alignment = Alignment(horizontal="right")

    for j in range(1, len(layout) + 1):
        width = max(len(str(c.value)) if c.value is not None else 0
                    for c in out[get_column_letter(j)])
        out.column_dimensions[get_column_letter(j)].width = min(max(width + 2, 10), 34)
    out.freeze_panes = "B2"
    out.auto_filter.ref = out.dimensions

    res = out_wb.create_sheet("Resumo")
    res.append(["Coluna higienizada", "Registros", "Marcados como inválidos"])
    for c in res[1]:
        c.font = Font(bold=True)
    for name, bad in stats.items():
        res.append([name, n, bad])
    res.column_dimensions["A"].width = 34
    res.column_dimensions["B"].width = 12
    res.column_dimensions["C"].width = 26
    out_wb.save(dst)
    return n, stats


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("entrada")
    ap.add_argument("saida")
    ap.add_argument("--sheet")
    a = ap.parse_args()
    total, st = run(a.entrada, a.saida, a.sheet)
    print(f"{total} registros processados -> {a.saida}")
    for k, v in st.items():
        print(f"  {k}: {v} inválidos")
